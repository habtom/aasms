"""PostgreSQL regression tests for the complete Excel export.

Set TEST_DATABASE_URL to a PostgreSQL database with the current Prisma schema.
Tests shadow its tables with temporary copies and roll back all fixture data.
"""

import os
from io import BytesIO

import pytest
from openpyxl import load_workbook
from sqlalchemy import create_engine, text
from sqlalchemy.orm import Session

from src.services.export_service import ExportService


@pytest.fixture
def export_db():
    database_url = os.environ.get("TEST_DATABASE_URL")
    if not database_url:
        pytest.skip("Set TEST_DATABASE_URL to run PostgreSQL export regression tests")

    engine = create_engine(database_url)
    try:
        with engine.connect() as connection, Session(bind=connection) as db:
            # LIKE keeps these tests aligned with the real schema, including the
            # absence of legacy columns. Temporary tables hide the originals.
            for table in (
                "Study",
                "ResearchQuestion",
                "StudyParameters",
                "InclusionCriterion",
                "ExclusionCriterion",
                "Facet",
                "FacetCategory",
                "Source",
                "SourceAnalysis",
                "Classification",
            ):
                db.execute(
                    text(
                        f'CREATE TEMP TABLE "{table}" (LIKE "{table}" INCLUDING ALL) ON COMMIT DROP'
                    )
                )

            db.execute(
                text("""
                INSERT INTO "Study" (id, title, "updatedAt") VALUES
                    ('study', 'Export regression study', CURRENT_TIMESTAMP),
                    ('empty-study', 'Empty study', CURRENT_TIMESTAMP),
                    ('other-study', 'Other study', CURRENT_TIMESTAMP)
            """)
            )
            yield db
            # Session.close() also rolls back if a test raises before this point.
            db.rollback()
    finally:
        engine.dispose()


@pytest.fixture
def populated_export_db(export_db):
    export_db.execute(
        text("""
        INSERT INTO "Facet" (id, "studyId", name, type, "order", "updatedAt") VALUES
            ('closed', 'study', 'Method', 'CLOSED', 0, CURRENT_TIMESTAMP),
            ('open', 'study', 'Notes', 'OPEN', 1, CURRENT_TIMESTAMP),
            ('coded', 'study', 'Theme', 'OPEN_CODED', 2, CURRENT_TIMESTAMP);
        INSERT INTO "FacetCategory" (id, "facetId", name, "order") VALUES
            ('survey', 'closed', 'Survey', 0),
            ('experiment', 'closed', 'Experiment', 1),
            ('safety', 'coded', 'Safety', 0);
        INSERT INTO "Source" (
            id, "studyId", type, "sourceCategory", title, authors, keywords,
            "publicationDate", "originalUrl", status, "finalDecision", "updatedAt"
        ) VALUES
            ('classified', 'study', 'PDF', 'FORMAL', 'A classified paper',
             ARRAY['Alice', 'Bob'], '{}', '2024-05-06', 'https://example.org/paper',
             'CLASSIFIED', 'INCLUDE', CURRENT_TIMESTAMP),
            ('no-analysis', 'study', 'PDF', 'FORMAL', 'B manually included',
             '{}', '{}', NULL, NULL, 'INCLUDED', 'INCLUDE', CURRENT_TIMESTAMP),
            ('no-classifications', 'study', 'PDF', 'FORMAL', 'C unclassified paper',
             '{}', '{}', '2023-01-01', NULL, 'INCLUDED', 'INCLUDE', CURRENT_TIMESTAMP),
            ('status-only', 'study', 'PDF', 'FORMAL', 'D classified by status',
             '{}', '{}', '2022-01-01', NULL, 'CLASSIFIED', NULL, CURRENT_TIMESTAMP),
            ('excluded', 'study', 'PDF', 'FORMAL', 'E excluded paper',
             ARRAY['Carol'], '{}', '2021-03-04', NULL, 'EXCLUDED', 'EXCLUDE', CURRENT_TIMESTAMP),
            ('other-source', 'other-study', 'PDF', 'FORMAL', 'Other study paper',
             '{}', '{}', '2020-01-01', NULL, 'CLASSIFIED', 'INCLUDE', CURRENT_TIMESTAMP);
        INSERT INTO "SourceAnalysis" (
            id, "sourceId", "extractedText", "inclusionRecommendation",
            "inclusionReasoning", "exclusionReasoning", "confidenceScore",
            "editedFields", "updatedAt"
        ) VALUES
            ('analysis', 'classified', '', true, 'Relevant', '', 0.95, '{}', CURRENT_TIMESTAMP),
            ('empty-analysis', 'no-classifications', '', true, 'Relevant', '', 0.8, '{}', CURRENT_TIMESTAMP),
            ('status-analysis', 'status-only', '', true, 'Relevant', '', 0.9, '{}', CURRENT_TIMESTAMP),
            ('excluded-analysis', 'excluded', '', false, '', 'Out of scope', 0.85, '{}', CURRENT_TIMESTAMP);
        INSERT INTO "Classification" (id, "analysisId", "facetId", "categoryId", value, confidence) VALUES
            ('c1', 'analysis', 'closed', 'survey', NULL, 0.9),
            ('c2', 'analysis', 'closed', 'experiment', NULL, 0.8),
            ('c3', 'analysis', 'open', NULL, 'Field observations', 0.9),
            ('c4', 'analysis', 'coded', 'safety', NULL, 0.9),
            ('c5', 'status-analysis', 'closed', 'experiment', NULL, 0.9)
    """)
    )
    return export_db


def test_full_export_uses_current_schema_and_preserves_classifications(populated_export_db):
    content = ExportService(populated_export_db).generate_full_export("study")
    wb = load_workbook(BytesIO(content))

    assert wb.sheetnames == [
        "Protocol",
        "All Sources",
        "Included Sources",
        "Exclusion Log",
        "Classification Matrix",
        "Data Dictionary",
    ]
    assert wb["Protocol"]["B1"].value == "Export regression study"

    all_sources = list(wb["All Sources"].values)
    assert len(all_sources) == 6  # Header and all five sources from this study only.
    assert [row[0] for row in all_sources[1:]] == [
        "classified",
        "no-analysis",
        "no-classifications",
        "status-only",
        "excluded",
    ]
    assert all_sources[1][2:4] == ("Alice; Bob", 2024)
    assert all_sources[1][7] == "https://example.org/paper"
    assert all_sources[2][3] is None  # Missing publication date stays blank.

    included = list(wb["Included Sources"].values)
    assert len(included) == 5
    assert included[0][6:] == ("Method", "Notes", "Theme")
    assert included[1][3] == 2024
    assert included[1][6:] == ("Survey; Experiment", "Field observations", "Safety")
    assert included[2][0] == "no-analysis"
    assert included[2][3] is None
    assert included[2][6:] == (None, None, None)
    assert included[3][0] == "no-classifications"
    assert included[3][6:] == (None, None, None)
    assert included[4][0] == "status-only"
    assert included[4][6:] == ("Experiment", None, None)

    excluded = list(wb["Exclusion Log"].values)
    assert excluded[1:] == [("E excluded paper", "Carol", 2021, "Out of scope", "85%")]

    matrix = list(wb["Classification Matrix"].values)
    assert matrix[0][2:] == ("Survey", "Experiment", "Safety")
    assert [row[2:] for row in matrix[1:5]] == [(1, 1, 1), (0, 0, 0), (0, 0, 0), (0, 1, 0)]
    assert matrix[5][2:] == ("=SUM(C2:C5)", "=SUM(D2:D5)", "=SUM(E2:E5)")
    assert [row[0] for row in list(wb["Data Dictionary"].values)[-3:]] == [
        "Method",
        "Notes",
        "Theme",
    ]


def test_full_export_handles_an_empty_study(export_db):
    content = ExportService(export_db).generate_full_export("empty-study")
    wb = load_workbook(BytesIO(content))

    assert len(wb.sheetnames) == 6
    assert wb["Protocol"]["B1"].value == "Empty study"
    for name in ("All Sources", "Included Sources", "Exclusion Log"):
        assert wb[name].max_row == 1
    assert list(wb["Classification Matrix"].values) == [("ID", "Title"), ("TOTAL", None)]


def test_full_export_rejects_a_missing_study(export_db):
    with pytest.raises(ValueError, match="Study missing not found"):
        ExportService(export_db).generate_full_export("missing")
